"""Test B: the two blind plans for each participant, generated from the hand-run Test A (D65, D68).

Input is Abram's filled-in results sheet (`validation/raw/results.xlsx`, never committed):
- the Swipes tab: item number and reaction 1-4;
- the Choices tab: task and A/B;
- the TestB tab: the pre-randomised label of the personalised plan;
- the Participants tab: launch channel, and the inspiring fields in column AG (Test A 5.7, optional, D77).

For each participant it fits δ_person from the swipes and choices (the same Laplace update the
app will use). It then plans the day with the shared brief, and writes:
- both plans in an identical format, as a WhatsApp message and a PDF, labelled A/B from the sheet;
- a review sheet for Abram only: why each stop was chosen, what was left out, and any legs to check.

The sent plans never say *why* a stop was chosen, because that would reveal which plan is
personalised and break the blind test.

    python -m brain.testb weights --results ../validation/raw/results.xlsx
    python -m brain.testb plans --results ../validation/raw/results.xlsx \\
        --brief ../validation/raw/testb/brief.yaml --generic ../validation/raw/testb/generic.yaml \\
        --out ../validation/raw/testb/out
    python -m brain.testb fields        # which aspiration fields have enough ready moments (D77)
"""

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from dataclasses import dataclass, field, replace
from pathlib import Path

import yaml

from .config import DATA_DIR
from .features import Conditions, conjoint_plan_features, moment_features
from .aspiration import AspirationMap, load_aspiration
from .moments import Moment, Point, load_moments, parse_windows
from .personal import DeltaPosterior
from .planner import (Candidate, DayPlan, LegWeights, candidate_for, extend_dwell, fits_day, hhmm, parse_hhmm,
                      plan_day, spread_to_end)
from .population import PopulationModel, load_population
from .scoring import conjoint_observation, contributions, named_weights, swipe_observation, update_person
from .travel import LondonEstimate

REACTIONS = {1: "loved", 2: "fine", 3: "didnt_enjoy", 4: "havent_been"}
FIRST_DATA_ROW = 6   # rows 1-3 notes, row 4 the example (ignored), row 5 headers


# ---- inputs -------------------------------------------------------------------------

@dataclass
class Participant:
    pid: str
    channel: str | None = None
    swipes: list[tuple[int, str]] = field(default_factory=list)   # (item number, reaction)
    choices: list[tuple[int, str]] = field(default_factory=list)  # (task, "A"/"B")
    personal_label: str | None = None                             # "A" or "B"
    fields_raw: str | None = None                                 # Test A 5.7, optional (D77)


def _rows(ws, ncols: int):
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, max_col=ncols, values_only=True):
        if row and row[0] not in (None, ""):
            yield row


def load_results(path: Path) -> dict[str, Participant]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    people: dict[str, Participant] = {}

    def get(pid) -> Participant:
        pid = str(pid).strip()
        return people.setdefault(pid, Participant(pid))

    for row in _rows(wb["Participants"], 33):
        p = get(row[0])
        p.channel = str(row[19]).strip() if row[19] else None
        p.fields_raw = str(row[32]).strip() if len(row) > 32 and row[32] not in (None, "") else None   # column AG
    for pid, item, _name, reaction, *_ in _rows(wb["Swipes"], 5):
        if item is not None and reaction is not None:
            get(pid).swipes.append((int(item), REACTIONS[int(reaction)]))
    for pid, task, choice in _rows(wb["Choices"], 3):
        if task is not None and choice:
            get(pid).choices.append((int(task), str(choice).strip().upper()))
    for pid, label, *_ in _rows(wb["TestB"], 2):
        if label:
            get(pid).personal_label = str(label).strip().upper()
    # participants with no Test A data yet are dropped
    return {k: v for k, v in people.items() if v.swipes or v.choices}


def load_swipe_items(path: Path) -> dict[int, Moment]:
    with open(path, encoding="utf-8") as f:
        d = yaml.safe_load(f)
    out = {}
    for num, it in d["items"].items():
        out[int(num)] = Moment.model_validate({
            "id": it.get("moment_id") or f"swipe_{int(num):02d}", "title": it["name"],
            "place": {"name": it["name"]}, "features": {"place": it["place"], "moment": it["moment"]},
            "cost_gbp": it.get("cost_gbp"), "source": "swipe-item"})
    return out


def load_swipe_cards(path: Path) -> dict[int, tuple[str, list[str]]]:
    """Each swipe place's name and the moment cards at the same place (D100)."""
    with open(path, encoding="utf-8") as f:
        d = yaml.safe_load(f)
    return {int(num): (it["name"], list(it.get("cards") or [])) for num, it in d["items"].items()}


def places_by_reaction(p: Participant, swipe_cards: dict[int, tuple[str, list[str]]]):
    """(loved, disliked): the swipe places, each with its card ids."""
    loved = [swipe_cards[i] for i, r in p.swipes if r == "loved" and i in swipe_cards]
    disliked = [swipe_cards[i] for i, r in p.swipes if r == "didnt_enjoy" and i in swipe_cards]
    return loved, disliked


def load_tasks(path: Path) -> dict[int, dict[str, dict]]:
    tasks: dict[int, dict[str, dict]] = {}
    with open(path, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            tasks.setdefault(int(r["task"]), {})[r["plan"]] = r
    return tasks


@dataclass
class Brief:
    """The same brief for both plans (test-plan.md): same day, times, number of stops and budget."""
    start_name: str
    start: Point
    start_time: str
    end_time: str
    stops: int
    budget_gbp: float | None = None
    day_label: str = "A day out in London"
    weekday: str | None = None     # "mon".."sun": cards closed that day are left out
    temp_c: float | None = None
    rain_prob: float | None = None
    time_limit_s: float = 2.0
    lunch_gbp: float = 15.0        # D102 (GUESS): the flexible lunch slot's cost per person


def load_brief(path: Path) -> Brief:
    with open(path, encoding="utf-8") as f:
        d = yaml.safe_load(f)
    d["start"] = Point(**d["start"])
    return Brief(**d)


# ---- fitting ------------------------------------------------------------------------

def person_fields(amap: AspirationMap | None, p: Participant) -> frozenset[str]:
    """The fields a participant picked (Test A 5.7). Unknown entries are reported, not guessed."""
    if amap is None or not p.fields_raw:
        return frozenset()
    fields, unknown = amap.parse(p.fields_raw)
    if unknown:
        print(f"{p.pid}: unknown inspiring field(s) {unknown} in Participants column AG; ignored", file=sys.stderr)
    return fields


def fit_participant(pop: PopulationModel, p: Participant, swipe_items: dict[int, Moment],
                    tasks: dict[int, dict[str, dict]], amap: AspirationMap | None = None) -> DeltaPosterior:
    fields = person_fields(amap, p)
    obs = []
    for item, reaction in p.swipes:
        m = swipe_items.get(item)
        asp = amap.match(m, fields) if (amap and m is not None) else 0.0
        if m is not None and (o := swipe_observation(pop, m, reaction, asp)):
            obs.append(o)
    for task, choice in p.choices:
        if task in tasks and choice in ("A", "B"):
            obs.append(conjoint_observation(pop, conjoint_plan_features(tasks[task]["A"]),
                                            conjoint_plan_features(tasks[task]["B"]), chose_a=choice == "A"))
    return update_person(pop, obs)


def readback_lines(pop: PopulationModel, post: DeltaPosterior, top: int = 6,
                   field_labels: list[str] | None = None) -> list[str]:
    """Confident weights only, as plain lines Abram can turn into the read-back message.

    The aspiration weight is left out of that list: its prior alone looks "confident", and it only
    matters for people who picked a field, so it is read back from their answer instead (D77).
    """
    ws = [w for w in named_weights(pop, post) if w.confident and w.name != "aspiration_match"]
    ws.sort(key=lambda w: abs(w.weight), reverse=True)
    lines = []
    if field_labels:
        asp = next(w for w in named_weights(pop, post) if w.name == "aspiration_match")
        lines.append(f"Said these fields inspire them: {', '.join(field_labels)} "
                     f"(moments tied to them get weight {asp.weight:+.2f})")
    for w in ws[:top]:
        if w.direction in ("stronger", "weaker"):   # sign-constrained, e.g. cost (D88)
            more = "more" if w.direction == "stronger" else "less"
            lines.append(f"{w.label}: seems to matter {more} to them than to most people "
                         f"(weight {w.weight:+.2f}; population {w.population:+.2f})")
            continue
        verb = "seems to put them off" if w.direction == "dislikes" else "seems to appeal to them"
        # A weight whose sd has barely moved from the prior is the population's view, not theirs (e.g. rain, D95)
        prior_only = w.sd is not None and w.sd >= 0.9 * pop.spec(w.name).delta_sd
        note = " (mostly the population prior, not their answers)" if prior_only else ""
        lines.append(f"{w.label}: {verb} (weight {w.weight:+.2f}, P(negative) {w.p_negative:.2f}){note}")
    if len(lines) == (1 if field_labels else 0):
        lines.append("No confident weights yet: use the persona and the open answers for the read-back.")
    return lines


# ---- planning -----------------------------------------------------------------------

@dataclass
class PlanNotes:
    """For the review only (never sent): the familiar favourite and what happened to each loved place (D100),
    the person's pace and the lunch stop (D102)."""
    favourite: str | None = None
    loves: list[tuple[str, str]] = field(default_factory=list)       # (place, what happened to it)
    dislikes: list[tuple[str, str]] = field(default_factory=list)
    pace: str = "even"                                               # relaxed | even | packed
    lunch: str = ""


BUFFER_WALK_MIN, BUFFER_TRANSIT_MIN = 5, 10   # D100 (GUESS): breathing space after a walk / a Tube or bus leg
DAY_NAMES = dict(mon="Mondays", tue="Tuesdays", wed="Wednesdays", thu="Thursdays", fri="Fridays",
                 sat="Saturdays", sun="Sundays")
# D102 (GUESSES): every plan that spans midday has one meal stop, starting 11:30-13:30: a food card, or a flexible
# "lunch near X" slot of 45 min. The bonus (in utility units) makes the solver keep a lunch over an extra sight.
LUNCH_START = (11 * 60 + 30, 13 * 60 + 30)
LUNCH_MIN = 45
LUNCH_BONUS = 5.0
LUNCH_CARD_MIN = 0.5   # a food card is offered as the lunch stop only if the person is predicted to like it this much


def _base(cid: str) -> str:
    """A card's id from a planner id: "m_x#lunch" (the card as the lunch stop) → "m_x"."""
    return cid.split("#")[0]


def pace_of(pop: PopulationModel, post: DeltaPosterior) -> str:
    """relaxed / packed from a confident packed-day weight, else even (D102)."""
    w = next((w for w in named_weights(pop, post) if w.name == "packed_pace"), None)
    if w is None or not w.confident:
        return "even"
    return "relaxed" if w.direction == "dislikes" else "packed"


def dwell_targets(by_id: dict[str, Moment], ids: list[str]) -> dict[str, int]:
    """D102: a relaxed person's stay at each stop: the top of the card's range, and at least 45 min in a garden
    or 60 in a park (a green place with "park" in its name)."""
    out = {}
    for mid in ids:
        m = by_id[mid]
        exp = m.experience.duration_min
        top = float(max(exp) if isinstance(exp, list) else exp)
        if (m.features.place.get("green_nature") or 0) >= 2:
            top = max(top, 60.0 if "park" in m.place.name.lower() else 45.0)
        out[mid] = int(round(float(m.approach.duration_min or 0) + top + float(m.experience.queue_min or 0)))
    return out


def lunch_candidates(real: list, by_id: dict[str, Moment], brief: Brief) -> list:
    """The planner's candidates plus the lunch options (D102): each food card the person is predicted to like
    (utility ≥ LUNCH_CARD_MIN) again as "<id>#lunch", allowed to start only in the lunch window (sharing a variant
    group with the card, so a plan holds one of the two), and a flexible "lunch near X" slot beside every place,
    scored as neutral (0). Nothing is added if the day doesn't span lunch."""
    lo, hi = LUNCH_START
    day_start, day_end = parse_hhmm(brief.start_time), parse_hhmm(brief.end_time)
    if day_end < lo + LUNCH_MIN or day_start > hi:
        return list(real)
    out = []
    for c in real:
        m = by_id[c.moment_id]
        if (m.features.place.get("food") or 0) >= 2 and (m.cost_gbp or 0) > 0 and c.utility >= LUNCH_CARD_MIN:
            wins = tuple((max(a, lo), min(b, hi + c.duration_min)) for a, b in c.windows
                         if max(a, lo) + c.duration_min <= min(b, hi + c.duration_min))
            if wins:
                c = replace(c, group=c.group or f"meal:{c.moment_id}")
                out.append(replace(c, moment_id=f"{c.moment_id}#lunch", windows=wins, lunch=True))
        out.append(c)
    out += [Candidate(f"lunch@{c.moment_id}", c.end, c.end, LUNCH_MIN, ((lo, hi + LUNCH_MIN),), 0.0,
                      float(brief.lunch_gbp), None, "moment", lunch=True) for c in real]
    return out


def leg_text(t: int | None, mode: str, came_from: str, exit_station: str | None, arrive_at: str | None,
             first: bool) -> str:
    """D102: a Tube or bus leg from a stop starts with the walk to a station, then names where you go."""
    if t is None or t <= 2:
        return ""
    if mode == "walk":
        return f"Getting there: about {t} min on foot from {came_from}."
    dest = f" to {arrive_at}" if arrive_at else ""
    if first:
        return f"Getting there: about {t} min by Tube or bus from {came_from}{dest}."
    walk = f"walk to {exit_station}" if exit_station else "walk to the nearest station"
    return f"Getting there: about {t} min: {walk}, then take the Tube or bus{dest}."


def personal_plan(pop: PopulationModel, post: DeltaPosterior, moments: list[Moment], brief: Brief,
                  travel: LondonEstimate, must_do: str | None = None, amap: AspirationMap | None = None,
                  fields: frozenset[str] = frozenset(), end_window: tuple[int, int] | None = None,
                  loved: list[tuple[str, list[str]]] = (), disliked: list[tuple[str, list[str]]] = (),
                  familiar: bool = False) -> tuple[DayPlan, list[dict], list[dict], PlanNotes]:
    """The personalised plan. `end_window` (earliest, latest end in minutes) keeps it ending within
    30 min of the generic plan (D97); if that can't be met, it plans without it and says so in `status`.

    D100: cards at places the person didn't enjoy are never planned. With `familiar`, the plan holds exactly one
    card at a place they loved (the best combination of the top three that fit the day), and the other stops are
    places they haven't marked as loved. `loved` and `disliked` are (swipe place, its card ids).
    D102: one of the stops is lunch when the day spans midday; a relaxed person gets longer stays where the day
    has room; rain swaps respect the person's weights."""
    b = pop.b_person(post.mean)
    cond = Conditions(brief.temp_c, brief.rain_prob)
    rain = brief.rain_prob is not None and brief.rain_prob >= 0.5
    usable = [m for m in moments if m.plannable() and m.open_on(brief.weekday)]
    if rain:
        usable = [m for m in usable if m.best_time.weather != ["dry"]]
    by_id = {m.id: m for m in usable}
    asp = {m.id: (amap.match(m, fields) if amap else 0.0) for m in usable}
    disliked_ids = {cid for _, ids in disliked for cid in ids} - {must_do}
    loved_ids = {cid for _, ids in loved for cid in ids} - disliked_ids
    floor = pop.cutpoints["rating"][1]
    real = [candidate_for(m, float(pop.vector(moment_features(m, cond, aspiration_match=asp[m.id])) @ b))
            for m in usable if m.id not in disliked_ids]
    cands = lunch_candidates(real, by_id, brief)
    by_cand = {c.moment_id: c for c in cands}
    idx = {n: i for i, n in enumerate(pop.names)}
    legs = LegWeights(b[idx["transfer_15_30"]], b[idx["transfer_30_plus"]], b[idx["packed_pace"]])
    day_start, day_end = parse_hhmm(brief.start_time), parse_hhmm(brief.end_time)
    common = dict(floor=floor, max_stops=brief.stops, min_stops=brief.stops,
                  budget_gbp=brief.budget_gbp, time_limit_s=brief.time_limit_s,
                  buffer_walk_min=BUFFER_WALK_MIN, buffer_transit_min=BUFFER_TRANSIT_MIN, lunch_bonus=LUNCH_BONUS)
    base_must = {must_do} if must_do else set()
    notes = PlanNotes(pace=pace_of(pop, post))
    targets = dwell_targets(by_id, [c.moment_id for c in real]) if notes.pace == "relaxed" else {}

    def solve(sub: list, must: set[str]) -> DayPlan:
        latest = min(day_end, end_window[1]) if end_window else day_end
        plan = None
        if end_window:   # D97: end within 30 min of the generic plan
            plan = plan_day(sub, brief.start, day_start, latest, travel, legs, must_include=must, **common)
            if plan.status != "ok" or not plan.stops:
                plan = None
        if plan is None:
            plan = plan_day(sub, brief.start, day_start, day_end, travel, legs, must_include=must, **common)
        if targets and plan.stops:
            extend_dwell(plan, by_cand, targets, latest)
        if end_window and plan.stops:
            ok = spread_to_end(plan, by_cand, end_window[0], latest)
            if (not ok or plan.stops[-1].depart_min > end_window[1]) and plan.status == "ok":
                plan.status = "could not end within 30 min of the generic plan: fix by hand"
        return plan

    loved_all = {c.moment_id for c in cands if _base(c.moment_id) in loved_ids}
    if familiar and loved_ids:
        options = sorted((c for c in real if c.moment_id in loved_ids and fits_day(c, day_start, day_end)
                          and (brief.budget_gbp is None or c.cost_gbp <= brief.budget_gbp)),
                         key=lambda c: c.utility, reverse=True)[:3]
        best = None
        for fav in options:
            forced_ids = [fav.moment_id] + [x for x in (f"{fav.moment_id}#lunch",) if x in by_cand]
            for forced in forced_ids:
                sub = [c for c in cands if c.moment_id not in loved_all or _base(c.moment_id) == fav.moment_id]
                trial = solve(sub, base_must | {forced})
                if trial.status == "ok" and forced in {st.moment_id for st in trial.stops} \
                        and (best is None or trial.value > best.value):
                    best, notes.favourite = trial, fav.moment_id
        plan = best if best is not None else solve([c for c in cands if c.moment_id not in loved_all], base_must)
    else:
        plan = solve(cands, base_must)

    stops, prev_point, prev_name, prev_exit = [], brief.start, brief.start_name, brief.start_name
    for st in plan.stops:
        c = by_cand[st.moment_id]
        mode = travel.mode(prev_point, c.start)
        if st.moment_id.startswith("lunch@"):
            host = by_id[st.moment_id[len("lunch@"):]]
            # the slot sits where the host card ends: after the host, that's where you are; otherwise (no end point
            # of its own) it's where you arrive for it, e.g. Cutty Sark DLR station before the Greenwich hill
            after_host = bool(stops) and _base(stops[-1]["moment_id"]) == host.id
            place = host.leave_name() if after_host or host.experience.end_point else \
                (clean(host.approach.from_) or host.leave_name())
            stop = {"title": f"Lunch near {place}", "approach": "",
                    "what": "Pick a café, pub or food stall nearby.", "best": "", "best_windows": [],
                    "queue": 0, "moment_id": st.moment_id, "predicted": None, "why": [],
                    "travel_text": leg_text(st.travel_min, mode, prev_name, prev_exit, host.approach.from_, not stops)}
            nxt = (c.end, place, host.exit_station())
        else:
            m = by_id[_base(st.moment_id)]
            stop = {"title": m.display_title(), "approach": approach_text(m, walked_from_stop=True),
                    "what": clean(m.experience.what_happens), "best": clean(m.best_time.light),
                    "best_windows": m.best_windows(), "queue": m.experience.queue_min or 0,
                    "moment_id": m.id, "predicted": st.utility,
                    "why": contributions(pop, b, moment_features(m, cond, st.travel_min, asp[m.id])),
                    "travel_text": leg_text(st.travel_min, mode, prev_name, prev_exit, m.approach.from_, not stops)}
            nxt = (m.end(), m.leave_name(), m.exit_station())
        stop.update({"arrive": hhmm(st.arrive_min), "depart": hhmm(st.depart_min), "travel_min": st.travel_min,
                     "travel_mode": mode, "from": prev_name, "cost_gbp": c.cost_gbp,
                     "cost_kind": "food" if c.lunch else by_id[_base(st.moment_id)].price_kind(),
                     "rain_swap": "", "lunch": c.lunch, "familiar": _base(st.moment_id) == notes.favourite})
        stops.append(stop)
        prev_point, prev_name, prev_exit = nxt
    lunch_stop = next((s for s in stops if s["lunch"]), None)
    notes.lunch = (f"{lunch_stop['title']} at {lunch_stop['arrive']}" if lunch_stop else
                   "none fitted: add one by hand" if len(cands) > len(real) else "the day doesn't span lunch")
    # Rain swaps are backups, not part of the planned day: a loved place can be one; a disliked place never is,
    # nor a place the person is predicted not to enjoy, nor one with a feature they clearly dislike (D102).
    dislikes = {w.name for w in named_weights(pop, post) if w.confident and w.direction == "dislikes"}
    allowed = {c.moment_id for c in real if c.utility >= floor
               and not any(moment_features(by_id[c.moment_id], cond, aspiration_match=asp[c.moment_id]).get(f, 0) >= 1.0
                           for f in dislikes)}
    add_rain_swaps(stops, plan, by_id, by_cand, allowed, brief, travel)
    # Only moments that fit the day's hours count as "left out" (D99: no evening cards in a 12:00-17:00 day)
    visited = {_base(s["moment_id"]) for s in stops}
    left_out = [{"moment_id": c.moment_id, "title": by_id[c.moment_id].title, "predicted": round(c.utility, 3)}
                for c in sorted(real, key=lambda c: c.utility, reverse=True)
                if fits_day(c, day_start, day_end) and c.moment_id not in visited][:3]
    all_ready = {m.id: m for m in moments}
    notes.loves = [(name, love_note(ids, all_ready, by_cand, plan, notes.favourite, familiar, brief, rain, travel))
                   for name, ids in loved]
    notes.dislikes = [(name, "kept out of the plan" if any(i in all_ready for i in ids) else "no moment card")
                      for name, ids in disliked]
    return plan, stops, left_out, notes


def love_note(ids: list[str], ready: dict[str, Moment], cands: dict, plan: DayPlan, favourite: str | None,
              familiar: bool, brief: Brief, rain: bool, travel: LondonEstimate) -> str:
    """Why a place the person loved is in the plan or not (D100): no card, closed, hours, budget, distance or the
    one-favourite rule. For Abram's review; never sent."""
    have = [i for i in ids if i in ready]
    if not have:
        return "no moment card yet"
    planned = {_base(st.moment_id) for st in plan.stops}
    if favourite in have:
        return "in the plan, as the familiar favourite"
    if planned & set(have):
        return "in the plan"
    day_start, day_end = parse_hhmm(brief.start_time), parse_hhmm(brief.end_time)

    def could_fit(m: Moment) -> bool:
        return (m.plannable() and m.open_on(brief.weekday) and not (rain and m.best_time.weather == ["dry"])
                and fits_day(candidate_for(m, 0.0), day_start, day_end)
                and (brief.budget_gbp is None or (m.cost_gbp or 0) <= brief.budget_gbp))

    # judge the place by its card that could fit the day, if any (the day boat, not the night boat)
    m = max((ready[i] for i in have), key=lambda m: (could_fit(m), cands[m.id].utility if m.id in cands else -99))
    if not m.plannable():
        why = "its card isn't complete yet"
    elif not m.open_on(brief.weekday):
        why = f"closed on {DAY_NAMES.get((brief.weekday or '')[:3], brief.weekday)}"
    elif rain and m.best_time.weather == ["dry"]:
        why = "it's outdoors and rain is forecast"
    elif not fits_day(candidate_for(m, 0.0), day_start, day_end):
        hours = ", ".join(f"{hhmm(a)}–{hhmm(b)}" for a, b in m.windows())
        why = f"its hours ({hours}) don't fit {brief.start_time}–{brief.end_time}"
    elif brief.budget_gbp is not None and (m.cost_gbp or 0) > brief.budget_gbp:
        why = f"it costs {_money(m.cost_gbp)}, over the {_money(brief.budget_gbp)} budget"
    elif familiar and favourite:
        why = f"only one familiar favourite per plan, and {_lower_article(ready[favourite].title)} fitted the day better"
    elif familiar:
        why = "it didn't fit alongside the other stops"
    else:
        why = "it scored lower than the stops chosen, with the travel between them"
    near = [travel.minutes(ready[_base(st.moment_id)].end(), m.start()) for st in plan.stops
            if _base(st.moment_id) in ready and m.start() is not None]
    if near:
        why += f"; about {round(min(near))} min from the nearest stop"
    return why


def _lower_article(title: str) -> str:
    first, _, rest = title.partition(" ")
    return f"{first.lower()} {rest}" if first in ("A", "An", "The") else title


def _outdoors(m: Moment) -> bool:
    return (m.features.place.get("outdoor_exposure") or 0) >= 2 or m.best_time.weather == ["dry"]


def add_rain_swaps(stops: list[dict], plan: DayPlan, by_id: dict[str, Moment], cands: dict, allowed: set[str],
                   brief: Brief, travel: LondonEstimate, max_min: float = 20.0) -> None:
    """D100: for each outdoor stop, an indoor or covered moment nearby that fits the same time, written as
    "If it rains: …". Each swap is used once, and only `allowed` moments are used (D102: the person's weights).
    If no stop gets one and there is no indoor stop, the best allowed indoor moment within 30 min of any stop (at
    any time of the day) goes on that stop."""
    day_start, day_end = parse_hhmm(brief.start_time), parse_hhmm(brief.end_time)
    planned = {_base(st.moment_id) for st in plan.stops}
    groups = {by_id[i].variant_of for i in planned if i in by_id and by_id[i].variant_of}
    spend = sum(s["cost_gbp"] for s in stops)
    pool = [m for m in by_id.values() if m.id not in planned and m.id in allowed and m.id in cands
            and m.kind == "moment" and not _outdoors(m) and (m.variant_of is None or m.variant_of not in groups)]
    used: set[str] = set()

    def pick(st, stop: dict, limit: float, any_time: bool):
        here = by_id[_base(st.moment_id)].start()
        best = None
        for m in pool:
            if m.id in used:
                continue
            c = cands[m.id]
            lo, hi = (day_start, day_end) if any_time else (max(day_start, st.arrive_min - 30),
                                                             min(day_end, st.arrive_min + 30 + c.duration_min))
            if not fits_day(c, lo, hi):
                continue
            if brief.budget_gbp is not None and spend - stop["cost_gbp"] + (m.cost_gbp or 0) > brief.budget_gbp:
                continue
            t = travel.minutes(here, m.start())
            if t > limit:
                continue
            score = c.utility - 0.02 * t
            if best is None or score > best[0]:
                best = (score, m, t, travel.mode(here, m.start()))
        return best

    def line(m: Moment, t: float, mode: str) -> str:
        how = "on foot" if mode == "walk" else "by Tube or bus"
        return f"{m.display_title()}, about {max(5, math.ceil(t / 5) * 5)} min away {how}."

    real = [(st, s) for st, s in zip(plan.stops, stops) if _base(st.moment_id) in by_id]
    outdoor = [(st, s) for st, s in real if _outdoors(by_id[_base(st.moment_id)])]
    for st, s in outdoor:
        if (hit := pick(st, s, max_min, any_time=False)):
            used.add(hit[1].id)
            s["rain_swap"] = line(hit[1], hit[2], hit[3])
    indoor = [st for st, _ in real if (by_id[_base(st.moment_id)].features.place.get("outdoor_exposure") or 0) == 0]
    if outdoor and not used and not indoor:
        hits = [(h, s) for st, s in outdoor if (h := pick(st, s, 30.0, any_time=True))]
        if hits:
            hit, s = max(hits, key=lambda x: x[0][0])
            s["rain_swap"] = line(hit[1], hit[2], hit[3])


GENERIC_STOP_MIN = 45   # GUESS: used when a generic stop has no `depart` time
OTHER_LABEL = {"A": "B", "B": "A", "1": "2", "2": "1"}
NUMBERED = {"A": "1", "B": "2"}   # --numbered: "Plan 1" / "Plan 2" instead of A / B


def plan_end(stops: list[dict]) -> int | None:
    """When a plan finishes, in minutes: the last stop's departure (or arrival + 45 min for a generic stop)."""
    if not stops:
        return None
    last = stops[-1]
    if last.get("depart"):
        return parse_hhmm(last["depart"])
    return parse_hhmm(last["arrive"]) + GENERIC_STOP_MIN


_VERBS = {"walk", "take", "go", "enter", "leave", "cross", "follow", "climb", "step", "ride", "head", "turn",
          "come", "start"}


def _same_place(a: str, b: str) -> bool:
    a, b = a.lower().strip(), b.lower().strip()
    return bool(a and b) and (a in b or b in a)


def approach_text(m: Moment, walked_from_stop: bool = False, came_from: str | None = None) -> str:
    """How to arrive, as full sentences (D100): "From X station, walk ..." when the path starts with a verb, else
    "Start at X. <path>". The card's "from" is left out after a walk from the previous stop (you aren't coming
    from it), when the path already says where from, and when you are already coming from there."""
    path = _sentence(clean(m.approach.path))
    frm = clean(m.approach.from_)
    if (not frm or walked_from_stop or path.lower().startswith("from ")
            or (came_from and _same_place(came_from, frm))):
        return path
    if not path:
        return f"Start at {frm}."
    if path.split()[0].lower().strip(",:;") in _VERBS:
        return f"From {frm}, {path[0].lower()}{path[1:]}"
    return f"Start at {frm}. {path}"


def _sentence(text: str) -> str:
    return text + "." if text and text[-1] not in ".!?" else text


def generic_stops(path: Path, moments: dict[str, Moment]) -> list[dict]:
    """The generic plan, written once by Abram. Stops are moment ids or free text, same fields.
    D100 fields: `best_window` (when the `best` line is shown), `cost_kind` (food/entry/ticket/hire) and
    `rain_swap` (the "If it rains:" line for an outdoor stop). D102: `lunch: true` marks the meal stop and
    `queue_min` adds an "Allow about N min to queue." line."""
    with open(path, encoding="utf-8") as f:
        d = yaml.safe_load(f)
    out = []
    for s in d["stops"]:
        if "moment_id" in s:
            m = moments[s["moment_id"]]
            s = {"title": m.title, "what": m.experience.what_happens or "", "best": m.best_time.light or "",
                 "best_window": m.best_time.best_window, "cost_gbp": m.cost_gbp or 0.0, "cost_kind": m.price_kind(),
                 "approach": approach_text(m, came_from=s.get("from")), **s}
        out.append({"title": s["title"], "arrive": s["arrive"], "depart": s.get("depart"), "travel_min": s.get("travel_min"),
                    "travel_mode": s.get("travel_mode", "walk"), "from": s.get("from", ""),
                    "approach": _sentence(clean(s.get("approach"))), "what": _sentence(clean(s.get("what"))),
                    "best": clean(s.get("best")), "best_windows": parse_windows(s.get("best_window") or [], s["title"]),
                    "cost_gbp": s.get("cost_gbp") or 0.0, "cost_kind": s.get("cost_kind") or "entry",
                    "rain_swap": _sentence(clean(s.get("rain_swap"))), "lunch": bool(s.get("lunch")),
                    "queue": s.get("queue_min") or 0})
    return out


# ---- rendering (identical for both plans) ------------------------------------------

_MARKER = re.compile(r"\s*\((?:GUESS|CHECKED|FOUNDER|FRIEND|UNVERIFIED)\b[^)]*\)")

# D100: the transport estimate. Each Tube or bus leg is counted as one Zone 1 pay-as-you-go single, capped at the
# Zones 1-2 daily cap. Sources: TfL fare API (London Bridge -> Westminster, valid from 4 Jul 2026: £3.00 off-peak,
# £3.10 peak; read 2026-09-28) and TfL adult fares 2026 (daily cap £8.90; read 2026-09-27).
TUBE_FARE_GBP = 3.00
DAILY_CAP_GBP = 8.90


def clean(text: str | None) -> str:
    """Participant-facing text: drop the cards' internal notes, e.g. "(GUESS)" or "(CHECKED: ...)"."""
    return _MARKER.sub("", text or "").strip()


def _travel_line(s: dict) -> str:
    if "travel_text" in s:          # personal stops: worded in personal_plan (D102)
        return s["travel_text"]
    if s.get("travel_min") is None:
        return ""
    how = "on foot" if s["travel_mode"] == "walk" else "by Tube or bus"
    src = f" from {s['from']}" if s.get("from") else ""
    return f"Getting there: about {int(s['travel_min'])} min {how}{src}."


def _money(c: float) -> str:
    return f"£{c:,.0f}" if float(c).is_integer() else f"£{c:,.2f}"


def _cost(c: float) -> str:
    return "free" if not c else _money(c)


def _cost_line(s: dict) -> str:
    """Never call food "entry" (D100)."""
    c = s.get("cost_gbp") or 0.0
    if not c:
        return "Cost: free"
    return {"food": f"Food: about {_money(c)}", "ticket": f"Ticket: {_money(c)}",
            "hire": f"Hire: {_money(c)}"}.get(s.get("cost_kind") or "entry", f"Entry: {_money(c)}")


def transport_gbp(stops: list[dict]) -> float:
    legs = sum(1 for s in stops if s.get("travel_min") is not None and s.get("travel_mode") != "walk")
    return round(min(legs * TUBE_FARE_GBP, DAILY_CAP_GBP), 2)


def _leave(s: dict) -> str:
    return s.get("depart") or hhmm(parse_hhmm(s["arrive"]) + GENERIC_STOP_MIN)


def _best_line(s: dict) -> str:
    """"Best: …" only when the stop's own time matches (D100)."""
    if not s.get("best"):
        return ""
    a, b = parse_hhmm(s["arrive"]), parse_hhmm(_leave(s))
    return f"Best: {s['best']}" if any(lo < b and a < hi for lo, hi in s.get("best_windows") or []) else ""


def _header(label: str, brief: Brief, stops: list[dict]) -> list[str]:
    spend, fare = sum(s["cost_gbp"] for s in stops), transport_gbp(stops)
    money = (f"about {_money(spend)} for food and entry" if spend else "free entry") + \
            (f", plus about {_money(fare)} transport" if fare else ", all on foot")
    lunch = ", including lunch" if any(s.get("lunch") for s in stops) else ""
    return [f"Plan {label}: {brief.day_label}",
            f"Start: {brief.start_name} at {brief.start_time} · {len(stops)} stops{lunch}", f"Costs: {money}"]


def _stop_lines(s: dict) -> list[str]:
    parts = (_travel_line(s), s.get("approach"), s.get("what"),
             s.get("queue") and f"Allow about {int(s['queue'])} min to queue.", _best_line(s), _cost_line(s),
             s.get("rain_swap") and f"If it rains: {s['rain_swap']}")
    return [x for x in parts if x]


def plan_text(label: str, brief: Brief, stops: list[dict]) -> str:
    """A WhatsApp message (WhatsApp shows *text* as bold)."""
    title, *rest = _header(label, brief, stops)
    lines = [f"*{title}*", *rest, ""]
    for i, s in enumerate(stops, 1):
        lines.append(f"*{i}. {s['arrive']}–{_leave(s)}  {s['title']}*")
        lines += _stop_lines(s)
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def plan_pdf(label: str, brief: Brief, stops: list[dict], path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    st = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm, title=f"Plan {label}")
    title, *rest = _header(label, brief, stops)
    flow = [Paragraph(escape(title), st["Title"])] + [Paragraph(escape(x), st["Normal"]) for x in rest]
    flow.append(Spacer(1, 6 * mm))
    for i, s in enumerate(stops, 1):
        flow.append(Paragraph(escape(f"{i}. {s['arrive']}–{_leave(s)}  {s['title']}"), st["Heading3"]))
        flow += [Paragraph(escape(x), st["Normal"]) for x in _stop_lines(s)]
        flow.append(Spacer(1, 4 * mm))
    doc.build(flow)


def _words(text: str) -> list[str]:
    return re.findall(r"[a-z0-9']+", text.lower())


def text_checks(stops: list[dict], n: int = 5) -> list[str]:
    """For the review (D100): lines that aren't full sentences, and arrival and description lines that repeat
    each other (5 or more words in a row)."""
    out = []
    for i, s in enumerate(stops, 1):
        for key in ("approach", "what", "rain_swap"):
            x = s.get(key) or ""
            if x and x[-1] not in ".!?":
                out.append(f"stop {i}: no full stop at the end of the {key} line")
        a, w = _words(s.get("approach") or ""), _words(s.get("what") or "")
        grams = {tuple(a[k:k + n]) for k in range(len(a) - n + 1)}
        if any(tuple(w[k:k + n]) in grams for k in range(len(w) - n + 1)):
            out.append(f"stop {i}: the arrival and description lines repeat each other")
    return out


def review_text(p: Participant, post: DeltaPosterior, plan: DayPlan, stops: list[dict], left_out: list[dict],
                pop: PopulationModel, brief: Brief, must_do: str | None, intent: dict,
                field_labels: list[str] | None = None, notes: PlanNotes | None = None,
                generic: list[dict] | None = None, familiar: bool = False) -> str:
    personal, generic_label = (p.personal_label or "?"), OTHER_LABEL.get(p.personal_label or "", "?")
    notes = notes or PlanNotes()
    out = [f"# {p.pid}: review before sending (for Abram only)", "",
           f"- Channel: {p.channel or 'unknown'}. Personalised plan is **{personal}**, generic is **{generic_label}** (TestB tab).",
           f"- Data used: {len(p.swipes)} swipes, {len(p.choices)} choices ({post.n_obs} informative).",
           f"- Planner status: {plan.status}.",
           f"- Trip intent: {intent or 'none recorded'}. Must-do: {must_do or 'none'}.",
           f"- Inspiring fields (Test A 5.7): {', '.join(field_labels) if field_labels else 'none given'}.",
           f"- Pace (D102): {notes.pace}"
           + (" (longer stays where the day has room)" if notes.pace == "relaxed" else "") + ".",
           f"- Lunch (D102): {notes.lunch}.",
           f"- Familiar favourite (D100, {'on' if familiar else 'off'}): "
           + (next((s['title'] for s in stops if s.get('familiar')), notes.favourite) if notes.favourite
              else "none" if familiar else "rule not used"), "",
           "## Read-back lines (confident weights only)", ""]
    out += [f"- {l}" for l in readback_lines(pop, post, field_labels=field_labels)]
    out += ["", "## Personalised plan: why each stop", ""]
    for s in stops:
        flag = " **check this leg**" if s["travel_mode"] != "walk" or s["travel_min"] > 30 else ""
        if s["predicted"] is None:     # a flexible lunch slot
            out.append(f"- {s['arrive']}–{s['depart']} {s['title']} (flexible lunch slot, D102). "
                       f"Travel {s['travel_min']} min ({s['travel_mode']}){flag}.")
            continue
        why = "; ".join(f"{w['label']} {w['contribution']:+.2f}" for w in s["why"])
        fam = " **Familiar favourite.**" if s.get("familiar") else ""
        meal = " Lunch stop." if s.get("lunch") else ""
        out.append(f"- {s['arrive']}–{s['depart']} {s['title']} (predicted {s['predicted']:+.2f}): {why}. "
                   f"Travel {s['travel_min']} min ({s['travel_mode']}){flag}.{fam}{meal}")
    total = sum(s["cost_gbp"] for s in stops)
    out += ["", f"Food and entry: {_cost(total)}" + (f" against a budget of {_cost(brief.budget_gbp)}" if brief.budget_gbp else "")
            + f"; transport about {_cost(transport_gbp(stops))}.", "",
            "## Places they loved: in the plan or why not (D100)", ""]
    out += [f"- {name}: {why}" for name, why in notes.loves] or ["- none marked as loved"]
    outdoor_no_swap = [s["title"] for s in stops if not s.get("rain_swap")]
    if stops and not any(s.get("rain_swap") for s in stops):
        out.insert(out.index("## Places they loved: in the plan or why not (D100)") - 1,
                   "\n**Rain: no indoor option found near this plan. Add one by hand before sending (D100).**")
    out += ["", "## Places they didn't enjoy", ""]
    out += [f"- {name}: {why}" for name, why in notes.dislikes] or ["- none"]
    out += ["", "## Strong moments left out", ""]
    out += [f"- {m['title']} (predicted {m['predicted']:+.2f})" for m in left_out] or ["- none"]
    checks = [f"personalised {c}" for c in text_checks(stops)] + [f"generic {c}" for c in text_checks(generic or [])]
    out += ["", "## Text checks (D100)", ""] + ([f"- {c}" for c in checks] or ["- none found"])
    out += ["", "## Before sending", "",
            "- [ ] Every Tube/bus leg checked (the times are rough estimates)",
            "- [ ] Opening hours and prices still right on the day",
            "- [ ] Same number of stops and same format as the generic plan",
            "- [ ] Both plans end within 30 min of each other (D97)",
            "- [ ] Each outdoor stop has a rain swap in both plans, or the plan has an indoor stop (D100)",
            "- [ ] Both plans have a lunch stop around midday (D102)",
            "- [ ] Nothing in the text says why a stop was chosen (keeps the test blind)", ""]
    return "\n".join(out)


# ---- CLI ----------------------------------------------------------------------------

def _load_common(data_dir: Path):
    pop = load_population(data_dir / "priors.yaml")
    return (pop, load_swipe_items(data_dir / "swipe_items.yaml"), load_tasks(data_dir / "conjoint_design.csv"),
            load_moments(data_dir / "moments"))


def _load_aspiration(data_dir: Path) -> AspirationMap | None:
    path = data_dir / "aspiration.yaml"
    return load_aspiration(path) if path.exists() else None


def coverage_lines(amap: AspirationMap, moments: list[Moment]) -> list[str]:
    lines = []
    for r in amap.coverage(moments):
        flag = "  <-- FLAG: fewer than %d ready" % amap.min_ready if r["flag"] else ""
        lines.append(f"{r['label']:18s} ready {r['ready']} (in {r['ready_groups']} variant groups), "
                     f"drafts {r['drafts']}{flag}")
    return lines


def cmd_fields(args) -> int:
    """Which aspiration fields have enough checked moments (D77)."""
    data = Path(args.data)
    amap = _load_aspiration(data)
    if amap is None:
        print("No data/aspiration.yaml", file=sys.stderr)
        return 2
    for line in coverage_lines(amap, load_moments(data / "moments")):
        print(line)
    return 0


def cmd_weights(args) -> int:
    pop, items, tasks, _ = _load_common(Path(args.data))
    amap = _load_aspiration(Path(args.data))
    people = load_results(Path(args.results))
    for pid in sorted(people):
        if args.ids and pid not in args.ids:
            continue
        post = fit_participant(pop, people[pid], items, tasks, amap)
        labels = amap.labels(person_fields(amap, people[pid])) if amap else []
        print(f"{pid} ({people[pid].channel or 'channel?'}, {post.n_obs} informative answers)")
        for line in readback_lines(pop, post, field_labels=labels):
            print(f"  - {line}")
    return 0


def cmd_plans(args) -> int:
    pop, items, tasks, moments = _load_common(Path(args.data))
    # Test B plans go to real people, so only cards Abram has checked (status: ready) are used, unless
    # --include-drafts is given for a dry run. Drafts can hold Claude's GUESS values (D72).
    drafts_skipped = [m.id for m in moments if m.status != "ready" and m.plannable()]
    if not args.include_drafts:
        moments = [m for m in moments if m.status == "ready"]
        if drafts_skipped:
            print(f"Skipping {len(drafts_skipped)} draft moments (set status: ready once checked, "
                  f"or use --include-drafts for a dry run): {', '.join(drafts_skipped)}", file=sys.stderr)
    elif drafts_skipped:
        print(f"DRY RUN: using {len(drafts_skipped)} unchecked draft moments. Don't send these plans.", file=sys.stderr)
    plannable = [m for m in moments if m.plannable()]
    amap = _load_aspiration(Path(args.data))
    if amap:
        flagged = [l for l in coverage_lines(amap, moments) if "FLAG" in l]
        if flagged:
            print("Aspiration fields with too few ready moments (people who picked them get little from it):\n  "
                  + "\n  ".join(flagged), file=sys.stderr)
    brief = load_brief(Path(args.brief))
    if len(plannable) < brief.stops:
        drafts = {m.id: m.missing_for_planning()[:3] for m in moments if not m.plannable()}
        print(f"Only {len(plannable)} plannable moments; the brief needs at least {brief.stops}. "
              f"Fill in data/moments (Area 2 schema). Drafts: {drafts}", file=sys.stderr)
        return 2
    people = load_results(Path(args.results))
    intents = {}
    if args.intent:
        with open(args.intent, encoding="utf-8") as f:
            intents = yaml.safe_load(f) or {}
    gen = generic_stops(Path(args.generic), {m.id: m for m in moments})
    gen_end = plan_end(gen)
    if gen and not any(s["lunch"] for s in gen):
        print("Warning: the generic plan has no stop marked 'lunch: true'. Personal plans include lunch when the day "
              "spans midday (D102).", file=sys.stderr)
    if gen and gen[0].get("from") and not _same_place(gen[0]["from"], brief.start_name):
        print(f"Warning: the generic plan starts from {gen[0]['from']!r} and the brief from {brief.start_name!r}. "
              f"Use the same start in both (D100).", file=sys.stderr)
    if gen and not gen[-1].get("depart"):
        print(f"Warning: the generic plan's last stop has no 'depart'; assuming {GENERIC_STOP_MIN} min there.",
              file=sys.stderr)
    end_window = (gen_end - 30, gen_end + 30) if gen_end is not None else None   # D97: end within 30 min
    if len(gen) != brief.stops:
        print(f"Warning: the generic plan has {len(gen)} stops and the brief asks for {brief.stops}.", file=sys.stderr)
    for field_ in ("approach", "what"):
        missing = [s["title"] for s in gen if not s.get(field_)]
        if missing:
            print(f"Warning: generic stops without '{field_}' ({missing}). Personalised stops have it, "
                  f"so the difference could give the plans away.", file=sys.stderr)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    travel = LondonEstimate()
    swipe_cards = load_swipe_cards(Path(args.data) / "swipe_items.yaml")
    for pid in sorted(people):
        if args.ids and pid not in args.ids:
            continue
        p = people[pid]
        if p.personal_label not in ("A", "B"):
            print(f"{pid}: no A/B label on the TestB tab, skipped", file=sys.stderr)
            continue
        intent = intents.get(pid, {}) or {}
        must_do = intent.get("must_do")
        post = fit_participant(pop, p, items, tasks, amap)
        fields = person_fields(amap, p)
        labels = amap.labels(fields) if amap else []
        loved, disliked = places_by_reaction(p, swipe_cards)
        plan, stops, left_out, notes = personal_plan(pop, post, moments, brief, travel, must_do, amap, fields,
                                                     end_window, loved, disliked, familiar=args.familiar)
        if args.numbered:
            p = replace(p, personal_label=NUMBERED[p.personal_label])
        generic_label = OTHER_LABEL[p.personal_label]
        for label, st in ((p.personal_label, stops), (generic_label, gen)):
            (out / f"{pid}_plan_{label}.txt").write_text(plan_text(label, brief, st), encoding="utf-8")
            if not args.no_pdf:
                plan_pdf(label, brief, st, out / f"{pid}_plan_{label}.pdf")
        review = review_text(p, post, plan, stops, left_out, pop, brief, must_do, intent, labels,
                             notes=notes, generic=gen, familiar=args.familiar)
        p_end = plan_end(stops)
        if p_end is not None and gen_end is not None:
            ok = abs(p_end - gen_end) <= 30
            review = review.replace("## Read-back lines", f"- Ends: personalised {hhmm(p_end)}, generic {hhmm(gen_end)} "
                                    f"({'within 30 min' if ok else '**more than 30 min apart: fix by hand**'}, D97).\n\n"
                                    "## Read-back lines", 1)
        (out / f"{pid}_review.md").write_text(review, encoding="utf-8")
        print(f"{pid}: plans {min(p.personal_label, generic_label)} and {max(p.personal_label, generic_label)} written ({plan.status})")
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="python -m brain.testb", description=__doc__.split("\n\n")[0])
    ap.add_argument("--data", default=str(DATA_DIR), help="brain data folder (priors, swipe items, moments)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    w = sub.add_parser("weights", help="read-back lines for each participant")
    w.add_argument("--results", required=True)
    w.add_argument("--ids", type=lambda s: set(s.split(",")))
    sub.add_parser("fields", help="aspiration fields: ready moments per field, flagged below the minimum (D77)")
    pl = sub.add_parser("plans", help="write both plans and a review sheet for each participant")
    pl.add_argument("--results", required=True)
    pl.add_argument("--brief", required=True)
    pl.add_argument("--generic", required=True)
    pl.add_argument("--intent")
    pl.add_argument("--out", required=True)
    pl.add_argument("--ids", type=lambda s: set(s.split(",")))
    pl.add_argument("--no-pdf", action="store_true")
    pl.add_argument("--numbered", action="store_true", help='label the plans "Plan 1" / "Plan 2" instead of A / B')
    pl.add_argument("--familiar", action="store_true",
                    help="D100 proposal: include exactly one place the person loved, if one fits the day")
    pl.add_argument("--include-drafts", action="store_true",
                    help="dry run: also plan with draft moments (unchecked GUESS values); never send these")
    args = ap.parse_args(argv)
    return {"weights": cmd_weights, "plans": cmd_plans, "fields": cmd_fields}[args.cmd](args)


if __name__ == "__main__":
    raise SystemExit(main())
