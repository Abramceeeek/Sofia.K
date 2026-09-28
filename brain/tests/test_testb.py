"""Test B plan generation, run against a filled-in copy of the real results-sheet template."""

import re
import shutil
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from brain import testb
from brain.planner import Candidate, LegWeights, plan_day
from brain.travel import LondonEstimate, WalkingEstimate
from brain.moments import Point

from conftest import DATA

REPO = Path(__file__).resolve().parents[2]
START = Point(lat=51.5081, lon=-0.1281)


def fill_results(path: Path):
    """Two made-up participants in the real template: P01 (label A) and P02 (label B)."""
    wb = load_workbook(REPO / "validation/results-sheet.xlsx")
    ws = wb["Participants"]
    ws["T6"], ws["T7"] = "stranger", "friends & family"
    ws["AG6"] = "Food, nature"          # Test A 5.7 (D77); P02 skipped the question
    sw = wb["Swipes"]
    rows = [("P01", 11, "Greenwich view", 1), ("P01", 7, "Camden Market", 3), ("P01", 13, "Primrose Hill", 1),
            ("P01", 3, "Borough Market", 3), ("P01", 1, "Big Ben", 4),
            ("P02", 3, "Borough Market", 1), ("P02", 14, "Brick Lane", 1), ("P02", 6, "Hyde Park", 3)]
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            sw.cell(row=6 + i, column=1 + j, value=v)
    ch = wb["Choices"]
    for i, t in enumerate(range(1, 11)):
        ch.cell(row=6 + i, column=1, value="P01")
        ch.cell(row=6 + i, column=2, value=t)
        ch.cell(row=6 + i, column=3, value="A" if t % 2 else "B")
    tb = wb["TestB"]
    tb["B6"], tb["B7"] = "A", "B"   # overwrite the pre-randomised labels for a known test
    wb.save(path)


@pytest.fixture
def data_dir(tmp_path, test_moments):
    d = tmp_path / "data"
    (d / "moments").mkdir(parents=True)
    for f in ("priors.yaml", "swipe_items.yaml", "conjoint_design.csv", "aspiration.yaml"):
        shutil.copy(DATA / f, d / f)
    for m in test_moments:
        (d / "moments" / f"{m.id}.yaml").write_text(yaml.safe_dump(m.model_dump(by_alias=True)))
    return d


def test_load_results_reads_the_real_template(tmp_path):
    f = tmp_path / "results.xlsx"
    fill_results(f)
    people = testb.load_results(f)
    assert set(people) == {"P01", "P02"}           # the other pre-filled IDs have no Test A data
    p1 = people["P01"]
    assert p1.channel == "stranger" and p1.personal_label == "A"
    assert (11, "loved") in p1.swipes and (1, "havent_been") in p1.swipes
    assert len(p1.choices) == 10
    assert people["P02"].channel == "friends & family" and people["P02"].personal_label == "B"


def test_swipe_items_cover_all_16_and_use_known_tags(pop):
    items = testb.load_swipe_items(DATA / "swipe_items.yaml")
    assert sorted(items) == list(range(1, 17))
    assert items[12].id == "m_southken_nhm_reveal"
    for m in items.values():
        assert all(m.tag(t) is not None for t in ["crowds", "awe", "food", "green_nature"])


def test_weights_from_test_a_move_in_the_right_direction(pop, tmp_path):
    f = tmp_path / "results.xlsx"
    fill_results(f)
    p = testb.load_results(f)["P02"]
    items = testb.load_swipe_items(DATA / "swipe_items.yaml")
    post = testb.fit_participant(pop, p, items, testb.load_tasks(DATA / "conjoint_design.csv"))
    b = pop.b_person(post.mean)
    food = pop.names.index("food")
    assert b[food] > pop.spec("food").mean       # loved two food places


def test_planner_hits_exact_stops_budget_and_must_do(pop, test_moments):
    cands = [Candidate(m.id, m.start(), m.end(), int(m.duration_min()), tuple(m.windows()), 0.5, float(m.cost_gbp or 0))
             for m in test_moments]
    legs = LegWeights(-0.3, -0.7)
    plan = plan_day(cands, START, 600, 1080, WalkingEstimate(), legs, max_stops=3, min_stops=3,
                    budget_gbp=20, must_include={"t_park"}, time_limit_s=0.5)
    ids = [s.moment_id for s in plan.stops]
    assert len(ids) == 3 and "t_park" in ids
    assert "t_pricey" not in ids                 # £80 is over the £20 budget
    cost = {c.moment_id: c.cost_gbp for c in cands}
    assert sum(cost[i] for i in ids) <= 20


def test_planner_says_so_when_it_cannot_fit_the_stops(pop, test_moments):
    cands = [Candidate(m.id, m.start(), m.end(), int(m.duration_min()), tuple(m.windows()), 0.5)
             for m in test_moments]
    plan = plan_day(cands, START, 600, 700, WalkingEstimate(), LegWeights(-0.3, -0.7), max_stops=6, min_stops=6,
                    time_limit_s=0.3)
    assert plan.status.startswith("could not fit 6 stops") and len(plan.stops) < 6


def test_london_estimate_walks_short_legs_and_rides_long_ones():
    t = LondonEstimate()
    near, far = Point(lat=51.5081, lon=-0.1281), Point(lat=51.4769, lon=-0.0005)   # ~9 km apart
    assert t.mode(near, Point(lat=51.5100, lon=-0.1300)) == "walk"
    assert t.mode(near, far) == "transit"
    assert t.minutes(near, far) < WalkingEstimate().minutes(near, far)


def test_plans_command_writes_blind_identical_plans_and_a_review(tmp_path, data_dir):
    results = tmp_path / "results.xlsx"
    fill_results(results)
    brief = dict(yaml.safe_load((DATA / "testb" / "brief.example.yaml").read_text()))
    brief.update(stops=3, budget_gbp=40, time_limit_s=0.5)
    (tmp_path / "brief.yaml").write_text(yaml.safe_dump(brief))
    generic = yaml.safe_load((DATA / "testb" / "generic.example.yaml").read_text())
    generic["stops"] = generic["stops"][:3]
    (tmp_path / "generic.yaml").write_text(yaml.safe_dump(generic))
    (tmp_path / "intent.yaml").write_text(yaml.safe_dump({"P01": {"mode": "explore", "must_do": "t_view"}}))
    out = tmp_path / "out"
    rc = testb.main(["--data", str(data_dir), "plans", "--results", str(results), "--brief", str(tmp_path / "brief.yaml"),
                     "--generic", str(tmp_path / "generic.yaml"), "--intent", str(tmp_path / "intent.yaml"),
                     "--out", str(out)])
    assert rc == 0
    for pid in ("P01", "P02"):
        a = (out / f"{pid}_plan_A.txt").read_text()
        b = (out / f"{pid}_plan_B.txt").read_text()
        assert a.startswith("*Plan A:") and b.startswith("*Plan B:")
        for text in (a, b):   # blind: same shape, no reasons, no model words
            assert len(re.findall(r"^(?:Cost|Food|Entry|Ticket|Hire): ", text, re.M)) == 3
            assert re.search(r"^Start: .+ at 12:00 · 3 stops, including lunch$", text, re.M)   # D100, D102
            assert re.search(r"^\*1\. \d\d:\d\d–\d\d:\d\d  ", text, re.M)          # D100: time ranges
            assert "entry costs" not in text
            assert "predicted" not in text and "weight" not in text and "Chosen" not in text
        assert (out / f"{pid}_plan_A.pdf").read_bytes().startswith(b"%PDF")
        assert (out / f"{pid}_plan_B.pdf").stat().st_size > 1000
    p1_personal = (out / "P01_plan_A.txt").read_text()   # P01's personalised plan is A
    assert "Test moment t_view" in p1_personal            # the must-do stop
    review = (out / "P01_review.md").read_text()
    assert "Personalised plan is **A**" in review and "Read-back lines" in review
    assert "Inspiring fields (Test A 5.7): Food, Nature." in review
    assert "Said these fields inspire them: Food, Nature" in review
    assert "Inspiring fields (Test A 5.7): none given." in (out / "P02_review.md").read_text()
    assert "The National Gallery" in (out / "P02_plan_A.txt").read_text()   # P02's generic plan is A
    assert "Food: about £15" in (out / "P02_plan_A.txt").read_text()        # D100: food is never "entry"
    review = (out / "P01_review.md").read_text()
    assert "- Ends: personalised" in review and "(within 30 min" in review          # D97 pair rule
    assert "Both plans end within 30 min" in review
    assert (out / "overlap.md").exists() and (out / "P01_personal.json").exists()   # D104


def test_plans_command_refuses_without_enough_moments(tmp_path):
    d = tmp_path / "data"
    (d / "moments").mkdir(parents=True)
    for f in ("priors.yaml", "swipe_items.yaml", "conjoint_design.csv"):
        shutil.copy(DATA / f, d / f)
    shutil.copy(DATA / "moments" / "m_southken_nhm_reveal.yaml", d / "moments")
    results = tmp_path / "results.xlsx"
    fill_results(results)
    rc = testb.main(["--data", str(d), "plans", "--results", str(results),
                     "--brief", str(DATA / "testb" / "brief.example.yaml"),
                     "--generic", str(DATA / "testb" / "generic.example.yaml"), "--out", str(tmp_path / "o")])
    assert rc == 2


def test_weights_command_prints_readback(tmp_path, data_dir, capsys):
    results = tmp_path / "results.xlsx"
    fill_results(results)
    assert testb.main(["--data", str(data_dir), "weights", "--results", str(results), "--ids", "P01"]) == 0
    out = capsys.readouterr().out
    assert out.startswith("P01 (stranger, 14 informative answers)")


def test_draft_moments_stay_out_of_real_plans(tmp_path, data_dir, capsys):
    # Drafts can hold Claude's GUESS values, so real Test B plans use only checked (ready) moments.
    for f in (data_dir / "moments").glob("*.yaml"):
        doc = yaml.safe_load(f.read_text())
        doc["status"] = "draft"
        f.write_text(yaml.safe_dump(doc))
    results = tmp_path / "results.xlsx"
    fill_results(results)
    args = ["--data", str(data_dir), "plans", "--results", str(results),
            "--brief", str(DATA / "testb" / "brief.example.yaml"),
            "--generic", str(DATA / "testb" / "generic.example.yaml"), "--out", str(tmp_path / "o"), "--no-pdf"]
    assert testb.main(args) == 2
    assert "Skipping" in capsys.readouterr().err
    assert testb.main(args + ["--include-drafts"]) == 0
    assert "DRY RUN" in capsys.readouterr().err


def test_plan_text_drops_internal_notes_and_names_the_last_place():
    assert testb.clean("Any time; at dusk the dome is lit (GUESS)") == "Any time; at dusk the dome is lit"
    assert testb.clean("Sunset; late autumn skies are best (CHECKED)") == "Sunset; late autumn skies are best"
    assert testb.clean("A walk (CHECKED: official page) here") == "A walk here"
    assert testb.clean("Pelicans (since 1664)") == "Pelicans (since 1664)"      # ordinary brackets stay
    from brain.moments import load_moments
    ms = {m.id: m for m in load_moments(DATA / "moments")}
    assert ms["m_uber_boat_west_night"].leave_name() == "Westminster Pier"
    assert ms["m_st_dunstan_garden"].leave_name() == "St Dunstan-in-the-East"
    assert ms["m_park_break_olympic"].leave_name() == "Queen Elizabeth Olympic Park"
    # After a walk from the previous stop, the card's station isn't where you come from
    sd = ms["m_st_dunstan_garden"]
    assert testb.approach_text(sd) == "From Monument or Tower Hill station, walk down St Dunstan's Hill towards the river."
    walked = testb.approach_text(sd, walked_from_stop=True)
    assert not walked.startswith("From Monument") and walked.endswith(".")


def test_real_brief_lets_sunset_and_river_cards_in():
    """D99: the real Test B day runs 12:00-17:00, so the river and sunset cards can be planned; evening-only ones can't."""
    import yaml
    from brain.moments import load_moments
    from brain.planner import candidate_for, fits_day, parse_hhmm
    brief = yaml.safe_load((DATA / "testb" / "brief.example.yaml").read_text())
    lo, hi = parse_hhmm(brief["start_time"]), parse_hhmm(brief["end_time"])
    assert (brief["start_time"], brief["end_time"]) == ("12:00", "17:00")
    ms = {m.id: m for m in load_moments(DATA / "moments")}
    fits = {mid: fits_day(candidate_for(ms[mid], 0.0), lo, hi) for mid in
            ("m_to_river_villiers", "m_to_river_somerset", "m_alderman_stairs_sunset", "m_waterloo_bridge_sunset",
             "a_soho_chinatown_evening", "m_uber_boat_west_night")}
    assert all(fits[m] for m in ("m_to_river_villiers", "m_to_river_somerset", "m_alderman_stairs_sunset",
                                 "m_waterloo_bridge_sunset"))
    assert not fits["a_soho_chinatown_evening"] and not fits["m_uber_boat_west_night"]
    generic = yaml.safe_load((DATA / "testb" / "generic.example.yaml").read_text())
    ends = [parse_hhmm(s["depart"]) for s in generic["stops"]]
    assert all(lo <= parse_hhmm(s["arrive"]) for s in generic["stops"]) and max(ends) <= hi



def _stop(**kw):
    base = dict(title="X", arrive="10:00", depart="10:45", travel_min=10, travel_mode="walk", **{"from": "A"},
                approach="Walk up the hill.", what="You see the view.", best="", best_windows=[], cost_gbp=0.0,
                cost_kind="entry", rain_swap="")
    base.update(kw)
    return base


def test_costs_best_lines_and_transport_are_worded_as_friends_asked():
    """D100: food is never "entry"; "Best" only when the time matches; transport is estimated and capped."""
    assert testb._cost_line(_stop(cost_gbp=15, cost_kind="food")) == "Food: about £15"
    assert testb._cost_line(_stop(cost_gbp=6.4)) == "Entry: £6.40"
    assert testb._cost_line(_stop()) == "Cost: free"
    sunset = [(16 * 60, 17 * 60)]
    assert testb._best_line(_stop(best="Sunset", best_windows=sunset)) == ""
    assert testb._best_line(_stop(best="Sunset", best_windows=sunset, arrive="16:10", depart="16:50")) == "Best: Sunset"
    assert testb._best_line(_stop(best="Daylight")) == ""          # no window: never shown
    legs = [_stop(travel_mode="transit") for _ in range(4)]
    assert testb.transport_gbp(legs[:1]) == 3.0 and testb.transport_gbp(legs) == 8.9   # capped at the daily cap
    brief = testb.Brief(start_name="London Bridge station", start=START, start_time="10:00", end_time="15:00", stops=2)
    head = testb._header("1", brief, [_stop(cost_gbp=15, cost_kind="food"), _stop(travel_mode="transit")])
    assert head[1] == "Start: London Bridge station at 10:00 · 2 stops"
    assert head[2] == "Costs: about £15 for food and entry, plus about £3 transport"
    assert testb._header("1", brief, [_stop()])[2] == "Costs: free entry, all on foot"


def test_text_checks_flag_repeats_and_missing_full_stops():
    stops = [_stop(approach="Take the lift to floor 15", what="You take the lift to floor 15 and step out.")]
    checks = testb.text_checks(stops)
    assert any("full stop" in c for c in checks) and any("repeat" in c for c in checks)
    assert testb.text_checks([_stop()]) == []


def test_familiar_favourite_and_places_they_didnt_enjoy():
    """D100: with --familiar the plan holds exactly one loved place; a disliked place is never planned."""
    from brain.moments import load_moments
    from brain.population import load_population
    from brain.scoring import update_person
    pop = load_population(DATA / "priors.yaml")
    post = update_person(pop, [])   # the population prior: no answers
    moments = [m for m in load_moments(DATA / "moments") if m.status == "ready"]
    brief = testb.Brief(start_name="London Bridge station", start=Point(lat=51.5049, lon=-0.0863), start_time="10:00",
                        end_time="15:00", stops=4, budget_gbp=40, weekday="sat", time_limit_s=1.0)
    loved = [("Borough Market", ["m_borough_market_graze"]), ("Brick Lane for food", ["m_beigel_bake_late"]),
             ("Camden Market", [])]
    disliked = [("Tate Modern Turbine Hall", ["m_millennium_bridge_dome"])]   # a stand-in card to keep out
    plan, stops, _, notes = testb.personal_plan(pop, post, moments, brief, LondonEstimate(), loved=loved,
                                                disliked=disliked, familiar=True)
    ids = [s["moment_id"] for s in stops]
    assert sum(i in ("m_borough_market_graze", "m_beigel_bake_late") for i in ids) == 1
    assert notes.favourite in ids and "m_millennium_bridge_dome" not in ids
    why = dict(notes.loves)
    assert why["Camden Market"] == "no moment card yet"
    assert any(w.startswith("in the plan") for w in why.values())
    assert any(w.startswith("only one familiar favourite") for w in why.values())
    assert dict(notes.dislikes)["Tate Modern Turbine Hall"] == "kept out of the plan"
    for s in stops:   # D100: rain swaps are indoor or covered moments, never already in the plan
        if s["rain_swap"]:
            assert s["rain_swap"].split(",")[0] not in [x["title"] for x in stops]


def test_approach_text_is_full_sentences():
    from brain.moments import load_moments
    ms = {m.id: m for m in load_moments(DATA / "moments")}
    assert testb.approach_text(ms["m_borough_market_graze"], came_from="London Bridge station") == \
        "Walk into the market; it sits right beside the station."
    assert testb.approach_text(ms["a_covent_garden_soho"]).startswith("Start at Covent Garden Piazza. ")
    assert testb.approach_text(ms["m_greenwich_wolfe_view"]).startswith("From Cutty Sark DLR station, walk into the park")
    for m in ms.values():
        if m.status == "ready":
            assert testb.approach_text(m).endswith("."), m.id


def _dry_brief(**kw):
    base = dict(start_name="London Bridge station", start=Point(lat=51.5049, lon=-0.0863), start_time="10:00",
                end_time="15:00", stops=4, budget_gbp=40, weekday="sat", time_limit_s=1.0)
    base.update(kw)
    return testb.Brief(**base)


def _posterior(pop, **deltas):
    """A made-up, confident posterior: δ on the named personal features, sd 0.1."""
    import numpy as np
    from brain.personal import DeltaPosterior
    names = list(pop.personal_names)
    mean = np.zeros(len(names))
    for k, v in deltas.items():
        mean[names.index(k)] = v
    return DeltaPosterior(mean, np.eye(len(names)) * 100.0, 0, 30)


def test_every_plan_spanning_midday_has_one_lunch_stop():
    """D102: a food card the person likes, or a flexible "lunch near X" slot, starting 11:30-13:30."""
    from brain.moments import load_moments
    from brain.population import load_population
    from brain.scoring import update_person
    pop = load_population(DATA / "priors.yaml")
    moments = [m for m in load_moments(DATA / "moments") if m.status == "ready"]
    plan, stops, _, notes = testb.personal_plan(pop, update_person(pop, []), moments, _dry_brief(), LondonEstimate())
    lunches = [s for s in stops if s["lunch"]]
    assert len(stops) == 4 and len(lunches) == 1
    assert "11:30" <= lunches[0]["arrive"] <= "13:30" and lunches[0]["cost_kind"] == "food"
    head = testb._header("1", _dry_brief(), stops)
    assert head[1].endswith("4 stops, including lunch")
    # a day that doesn't span lunch gets none
    _, stops, _, notes = testb.personal_plan(pop, update_person(pop, []), moments,
                                            _dry_brief(start_time="15:00", end_time="20:00", stops=2), LondonEstimate())
    assert not any(s["lunch"] for s in stops) and notes.lunch == "the day doesn't span lunch"


def test_rain_swaps_respect_a_clear_dislike_of_crowds():
    """D102: someone who clearly dislikes crowds is never sent to a crowded market when it rains."""
    from brain.moments import load_moments
    from brain.population import load_population
    pop = load_population(DATA / "priors.yaml")
    moments = [m for m in load_moments(DATA / "moments") if m.status == "ready"]
    post = _posterior(pop, crowds=-2.0)
    assert any(w.name == "crowds" and w.direction == "dislikes" and w.confident
               for w in __import__("brain.scoring", fromlist=["x"]).named_weights(pop, post))
    _, stops, _, _ = testb.personal_plan(pop, post, moments, _dry_brief(), LondonEstimate())
    swaps = " ".join(s["rain_swap"] for s in stops)
    assert "Borough Market" not in swaps


def test_relaxed_people_stay_longer_at_gardens():
    """D102: pace sets dwell time: a relaxed person gets at least 45 min in a garden where the day allows."""
    from brain.moments import load_moments
    from brain.population import load_population
    from brain.planner import parse_hhmm
    pop = load_population(DATA / "priors.yaml")
    moments = [m for m in load_moments(DATA / "moments") if m.status == "ready"]
    relaxed = _posterior(pop, packed_pace=-3.0)
    assert testb.pace_of(pop, relaxed) == "relaxed"
    by_id = {m.id: m for m in moments}
    targets = testb.dwell_targets(by_id, ["m_st_dunstan_garden", "m_park_break_hyde", "m_horizon_22"])
    assert targets["m_st_dunstan_garden"] >= 45 + 5 and targets["m_park_break_hyde"] >= 60 + 8
    assert targets["m_horizon_22"] == 10 + 45          # not green: the top of its own range


def test_travel_legs_start_with_the_walk_to_a_station():
    """D102: "walk to X station, then take the Tube or bus to Y" when a leg starts at a stop."""
    t = testb.leg_text(22, "transit", "The Garden at 120", "Bank station", "Liverpool Street station", first=False)
    assert t == "Getting there: about 22 min: walk to Bank station, then take the Tube or bus to Liverpool Street station."
    assert testb.leg_text(20, "transit", "London Bridge station", "London Bridge station", "Cutty Sark DLR station",
                          first=True) == ("Getting there: about 20 min by Tube or bus from London Bridge station "
                                          "to Cutty Sark DLR station.")
    assert testb.leg_text(0, "walk", "X", None, None, first=False) == ""     # a lunch slot right where you are


def test_titles_name_the_place_and_queues_are_told():
    """D102: the place's name is in each title; a popular food stop says how long to allow for the queue."""
    from brain.moments import load_moments
    ms = {m.id: m for m in load_moments(DATA / "moments")}
    assert ms["m_garden_at_120"].display_title() == "The Garden at 120: a rooftop garden among the City towers"
    assert ms["m_beigel_bake_late"].display_title() == "Beigel Bake: a salt beef bagel on Brick Lane"
    assert ms["a_soho_chinatown_evening"].display_title() == ms["a_soho_chinatown_evening"].title   # already named
    assert ms["m_beigel_bake_late"].duration_min() == 10 + 22.5 + 10     # approach + experience + queue
    assert "Allow about 10 min to queue." in testb._stop_lines(_stop(queue=10))


def test_overlap_report_warns_when_dissimilar_people_share_most_stops(tmp_path):
    """D104: shared stops across participants; a warning when two dissimilar people share more than half."""
    import json
    def person(pid, ids, taste):
        stops = [{"id": i, "title": i.upper()} for i in ids]
        (tmp_path / f"{pid}_personal.json").write_text(json.dumps({"pid": pid, "stops": stops, "taste": taste}))
    person("P01", ["a", "b", "c", "lunch@a"], [2.0, 0.0, 0.0])
    person("P02", ["a", "b", "c", "d"], [-2.0, 0.0, 0.0])        # opposite tastes, 3 of 4 shared
    person("P03", ["a", "b", "e", "f"], [2.0, 0.1, 0.0])         # similar to P01, 2 of 4 shared
    person("P04", ["x", "y", "z", "w"], [0.0, 0.0, 0.1])         # too little data to compare
    report, warnings = testb.overlap_report(tmp_path)
    assert len(warnings) == 1 and warnings[0].startswith("P01 and P02 are dissimilar")
    assert "| P01–P03 | 2 of 4 | +1.00 |" in report
    assert "too little data" in report
    assert "- A: 3 of 4 plans" in report                        # a stop in most plans
    assert "LUNCH@A" not in report.split("## Pairs")[1].split("## Stops")[0]   # flexible lunch slots don't count


def test_meal_cards_are_only_ever_the_lunch_stop():
    """D104: "Lunch in the crypt…" is offered only as the lunch stop; and there is a lunch slot at the start."""
    from brain.moments import load_moments
    from brain.planner import candidate_for
    ms = {m.id: m for m in load_moments(DATA / "moments") if m.status == "ready"}
    real = [candidate_for(m, 1.0) for m in ms.values() if m.plannable()]
    ids = {c.moment_id for c in testb.lunch_candidates(real, ms, _dry_brief(start_time="12:00", end_time="17:00"))}
    assert "m_st_martin_crypt_lunch" not in ids and "m_st_martin_crypt_lunch#lunch" in ids
    assert "lunch@start" in ids
    evening = testb.lunch_candidates(real, ms, _dry_brief(start_time="15:00", end_time="20:00"))
    assert not any(c.moment_id.startswith("m_st_martin_crypt_lunch") for c in evening)
